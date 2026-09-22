
# ===== cell 0 =====
# ms_dmt_analysis_run5.py -- final consolidation. Builds BOTH:
#   (A) Primary Analysis (Uncapped): the full cohort, DMT data with NO date restriction
#       at all (full history, whenever records exist).
#   (B) Secondary Analysis (Decade): the same cohort, DMT data restricted to
#       2015-01-01..2026-06-30 (the "last decade" window already used throughout this
#       project), including year-by-year efficacy-tier switch-transition trends.
# Both analyses apply the SAME new exclusion criterion: patients with age at diagnosis
# >= 100 years are removed from the cohort entirely (not just flagged) -- this is a new
# rule; the original analyses retained these patients as "plausible, not capped".
#
# No earlier script is modified. Checkpoints/data are reused wherever already fetched:
#   - Cohort + unrestricted first-ever diagnosis date: ./_checkpoints/ms_cohort_demo.pkl
#   - Unrestricted (full-history) Medications DMT data: ./_checkpoints/ms_dmt_final.pkl
#   - 2015-2026-06 Medications+Procedures data: ./_checkpoints2/ + ./_checkpoints4/
#   - NEW: unrestricted (full-history) Procedures DMT data -- not previously pulled by
#     any prior script (every earlier Procedures pull was date-restricted) -- fetched
#     fresh here, batched, checkpointed to ./_checkpoints5/.
import os
import time

import numpy as np
import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL

# ===== cell 1 =====
load_dotenv("/Users/anjali28/Desktop/Eli lily AD model/.env")
db_host = os.getenv("DB_HOST")
db_port = os.getenv("DB_PORT", "3306")
db_user = os.getenv("DB_USER")
db_password = os.getenv("DB_PASSWORD")

url = URL.create(
    drivername="mysql+pymysql", username=db_user, password=db_password,
    host=db_host, port=int(db_port), database="rgd_udm_silver",
)
engine = create_engine(
    url, pool_pre_ping=True,
    connect_args={"connect_timeout": 30, "read_timeout": 7200, "write_timeout": 7200},
)

CHECKPOINT_DIR = "./_checkpoints5"
os.makedirs(CHECKPOINT_DIR, exist_ok=True)
DECADE_START = "2015-01-01"
DECADE_END = "2026-06-30"


def log(msg):
    print(f"[run5] {msg}", flush=True)


# ===== cell 2 =====
# ── Cohort + age cleanup (NEW: exclude age at diagnosis >= 100) ─────────────────────
ms_cohort_demo_raw = pd.read_pickle("./_checkpoints/ms_cohort_demo.pkl")
n_before = len(ms_cohort_demo_raw)
excluded_age100 = ms_cohort_demo_raw[ms_cohort_demo_raw["age_at_diagnosis"] >= 100]
ms_cohort_demo = ms_cohort_demo_raw[~(ms_cohort_demo_raw["age_at_diagnosis"] >= 100)].copy()
n_after = len(ms_cohort_demo)
log(f"age cleanup: {n_before:,} -> {n_after:,} patients ({len(excluded_age100)} excluded, age >= 100)")

cohort_ndids_clean = set(ms_cohort_demo["ndid"])
total_cohort_n = n_after


def _gender_bucket(g):
    if pd.isna(g):
        return "Unknown"
    g_low = str(g).strip().lower()
    if g_low.startswith("f"):
        return "Female"
    if g_low.startswith("m"):
        return "Male"
    return "Unknown"


ms_cohort_demo["gender_bucket"] = ms_cohort_demo["gender"].apply(_gender_bucket)


def _onset_tier_v2(age):
    if pd.isna(age):
        return "Unknown"
    if age < 18:
        return "Pediatric-Onset (<18)"
    if age == 18:
        return "Unclassified (18)"
    if age <= 49:
        return "Adult-Onset (19-49)"
    return "Late-Onset (50+)"


ms_cohort_demo["onset_tier_v2"] = ms_cohort_demo["age_at_diagnosis"].apply(_onset_tier_v2)

# ===== cell 3 =====
# ── DMT reference (unchanged) ────────────────────────────────────────────────────────
DMT_REFERENCE = [
    ("Interferon beta-1a",    "Interferon beta",               "Standard efficacy", "Injectable", ["interferon beta-1a", "interferon beta 1a", "avonex", "rebif"], ["peginterferon"]),
    ("Interferon beta-1b",    "Interferon beta",               "Standard efficacy", "Injectable", ["interferon beta-1b", "interferon beta 1b", "betaseron", "extavia"], []),
    ("Peginterferon beta-1a", "Interferon beta",               "Standard efficacy", "Injectable", ["peginterferon beta-1a", "peginterferon beta 1a", "plegridy"], []),
    ("Glatiramer acetate",    "Glatiramoid",                    "Standard efficacy", "Injectable", ["glatiramer", "copaxone", "glatopa"], []),
    ("Dimethyl fumarate",     "Fumarate",                       "Standard efficacy", "Oral",       ["dimethyl fumarate", "tecfidera"], []),
    ("Diroximel fumarate",    "Fumarate",                       "Standard efficacy", "Oral",       ["diroximel fumarate", "vumerity"], []),
    ("Monomethyl fumarate",   "Fumarate",                       "Standard efficacy", "Oral",       ["monomethyl fumarate", "bafiertam"], []),
    ("Teriflunomide",         "Pyrimidine synthesis inhibitor", "Standard efficacy", "Oral",       ["teriflunomide", "aubagio"], []),
    ("Fingolimod",            "S1P modulator",                  "High efficacy",     "Oral",       ["fingolimod", "gilenya"], []),
    ("Siponimod",             "S1P modulator",                  "High efficacy",     "Oral",       ["siponimod", "mayzent"], []),
    ("Ozanimod",              "S1P modulator",                  "High efficacy",     "Oral",       ["ozanimod", "zeposia"], []),
    ("Ponesimod",             "S1P modulator",                  "High efficacy",     "Oral",       ["ponesimod", "ponvory"], []),
    ("Cladribine",            "Purine analog",                  "High efficacy",     "Oral",       ["cladribine", "mavenclad"], []),
    ("Natalizumab",           "Anti-alpha4-integrin",           "High efficacy",     "Infusion",   ["natalizumab", "tysabri"], []),
    ("Alemtuzumab",           "Anti-CD52",                      "High efficacy",     "Infusion",   ["alemtuzumab", "lemtrada"], []),
    ("Ocrelizumab",           "Anti-CD20",                      "High efficacy",     "Infusion",   ["ocrelizumab", "ocrevus"], []),
    ("Ofatumumab",            "Anti-CD20",                      "High efficacy",     "Injectable", ["ofatumumab", "kesimpta"], []),
    ("Ublituximab",           "Anti-CD20",                      "High efficacy",     "Infusion",   ["ublituximab", "briumvi"], []),
    ("Mitoxantrone",          "Immunosuppressant",              "High efficacy",     "Infusion",   ["mitoxantrone", "novantrone"], []),
]
class_map = {g: c for g, c, _t, _r, _te, _ex in DMT_REFERENCE}
tier_map = {g: t for g, _c, t, _r, _te, _ex in DMT_REFERENCE}
all_classes = sorted({row[1] for row in DMT_REFERENCE})

HCPCS_DMT_REFERENCE = [
    ("J2350", "Ocrelizumab"), ("J2323", "Natalizumab"), ("J0202", "Alemtuzumab"),
    ("J1826", "Interferon beta-1a"), ("Q3027", "Interferon beta-1a"), ("J1595", "Glatiramer acetate"),
]
hcpcs_case_clauses = " ".join(
    f"WHEN proc_code_std = '{code}' OR proc_code = '{code}' THEN '{generic}'"
    for code, generic in HCPCS_DMT_REFERENCE
)
hcpcs_generic_case_sql = "CASE " + hcpcs_case_clauses + " END"
hcpcs_code_list_sql = ",".join(f"'{code}'" for code, _g in HCPCS_DMT_REFERENCE)

# ===== cell 4 =====
# ── NEW: unrestricted (full-history) Procedures pull, batched over the age-cleaned
# full cohort -- no prior script ever pulled Procedures without a date filter ───────
BATCH_SIZE = 1000
all_ndids_str = sorted(str(x) for x in cohort_ndids_clean)

PROCS_UNCAPPED_BATCH_DIR = os.path.join(CHECKPOINT_DIR, "procs_uncapped_batches")
os.makedirs(PROCS_UNCAPPED_BATCH_DIR, exist_ok=True)
PROCS_UNCAPPED_FINAL = os.path.join(CHECKPOINT_DIR, "procs_uncapped_final.pkl")

if os.path.exists(PROCS_UNCAPPED_FINAL):
    ms_dmt_procs_uncapped = pd.read_pickle(PROCS_UNCAPPED_FINAL)
    log(f"procs_uncapped: loaded {len(ms_dmt_procs_uncapped):,} rows from checkpoint")
else:
    t_start = time.time()
    n_batches = (len(all_ndids_str) + BATCH_SIZE - 1) // BATCH_SIZE
    for b in range(n_batches):
        batch_path = os.path.join(PROCS_UNCAPPED_BATCH_DIR, f"batch_{b:04d}.pkl")
        if os.path.exists(batch_path):
            continue
        i = b * BATCH_SIZE
        batch = all_ndids_str[i : i + BATCH_SIZE]
        id_list = ",".join(batch)
        sql = text(f"""
            SELECT
                ndid,
                {hcpcs_generic_case_sql}  AS drug_generic_name,
                proc_code, proc_code_std, proc_name,
                proc_start_date                                             AS proc_start_date_raw,
                encounter_date, enc_date_proxy,
                COALESCE(proc_start_date, encounter_date, enc_date_proxy)   AS med_start_date,
                ehr_source_name
            FROM rgd_udm_silver.procedures FORCE INDEX (idx_ndid)
            WHERE ndid IN ({id_list})
              AND udm_active_flag = 'Y'
              AND (proc_code_std IN ({hcpcs_code_list_sql}) OR proc_code IN ({hcpcs_code_list_sql}))
        """)
        last_exc = None
        for attempt in range(1, 5):
            try:
                with engine.connect() as conn:
                    batch_df = pd.read_sql(sql, conn)
                batch_df.to_pickle(batch_path)
                log(f"[procs-uncapped {i+len(batch)}/{len(all_ndids_str)}] batch {b} matches: "
                    f"{len(batch_df)} (elapsed {time.time()-t_start:.0f}s)")
                break
            except Exception as exc:
                last_exc = exc
                log(f"procs-uncapped batch {b} attempt {attempt} failed -- {exc}")
                if attempt < 4:
                    time.sleep(30 * attempt)
        else:
            raise last_exc
    results = [pd.read_pickle(os.path.join(PROCS_UNCAPPED_BATCH_DIR, f"batch_{b:04d}.pkl")) for b in range(n_batches)]
    ms_dmt_procs_uncapped = pd.concat(results, ignore_index=True) if results else pd.DataFrame()
    for _col in ["proc_start_date_raw", "encounter_date", "enc_date_proxy", "med_start_date"]:
        if _col in ms_dmt_procs_uncapped.columns:
            ms_dmt_procs_uncapped[_col] = pd.to_datetime(ms_dmt_procs_uncapped[_col], errors="coerce")
    if len(ms_dmt_procs_uncapped):
        ms_dmt_procs_uncapped["med_end_date"] = ms_dmt_procs_uncapped["med_start_date"]
        ms_dmt_procs_uncapped["drug_class"] = ms_dmt_procs_uncapped["drug_generic_name"].map(class_map)
        ms_dmt_procs_uncapped["efficacy_tier"] = ms_dmt_procs_uncapped["drug_generic_name"].map(tier_map)
        ms_dmt_procs_uncapped["discont_date"] = pd.NaT
        ms_dmt_procs_uncapped["discont_reason"] = None
    ms_dmt_procs_uncapped.to_pickle(PROCS_UNCAPPED_FINAL)

log(f"procs_uncapped: {len(ms_dmt_procs_uncapped):,} rows")

# ===== cell 5 =====
# ── Assemble ms_dmt_uncapped (Primary Analysis: NO date restriction) ────────────────
ms_dmt_meds_uncapped = pd.read_pickle("./_checkpoints/ms_dmt_final.pkl")
ms_dmt_meds_uncapped = ms_dmt_meds_uncapped[ms_dmt_meds_uncapped["ndid"].isin(cohort_ndids_clean)].copy()
ms_dmt_meds_uncapped["source_table"] = "Medications"
ms_dmt_procs_uncapped["source_table"] = "Procedures"

ms_dmt_uncapped = pd.concat([ms_dmt_meds_uncapped, ms_dmt_procs_uncapped], ignore_index=True, sort=False)
log(f"ms_dmt_uncapped: {len(ms_dmt_uncapped):,} rows, "
    f"{ms_dmt_uncapped['ndid'].nunique():,} treated patients (of {total_cohort_n:,})")

# ===== cell 6 =====
# ── Assemble ms_dmt_decade (Secondary Analysis: 2015-01-01..2026-06-30) -- reuse run2
# + run4's already date-restricted pulls, which together cover all 51,199 patients ──
meds_r2 = pd.read_pickle("./_checkpoints2/ms_dmt_meds_final.pkl")
procs_r2 = pd.read_pickle("./_checkpoints2/ms_dmt_procs_final.pkl")
meds_rem = pd.read_pickle("./_checkpoints4/meds_remaining_final.pkl")
procs_rem = pd.read_pickle("./_checkpoints4/procs_remaining_final.pkl")

ms_dmt_meds_decade = pd.concat([meds_r2, meds_rem], ignore_index=True, sort=False)
ms_dmt_meds_decade = ms_dmt_meds_decade[ms_dmt_meds_decade["ndid"].isin(cohort_ndids_clean)].copy()
ms_dmt_meds_decade["source_table"] = "Medications"

ms_dmt_procs_decade = pd.concat([procs_r2, procs_rem], ignore_index=True, sort=False)
ms_dmt_procs_decade = ms_dmt_procs_decade[ms_dmt_procs_decade["ndid"].isin(cohort_ndids_clean)].copy()
if "med_end_date" not in ms_dmt_procs_decade.columns:
    ms_dmt_procs_decade["med_end_date"] = ms_dmt_procs_decade["med_start_date"]
if "discont_date" not in ms_dmt_procs_decade.columns:
    ms_dmt_procs_decade["discont_date"] = pd.NaT
if "discont_reason" not in ms_dmt_procs_decade.columns:
    ms_dmt_procs_decade["discont_reason"] = None
ms_dmt_procs_decade["drug_class"] = ms_dmt_procs_decade["drug_generic_name"].map(class_map)
ms_dmt_procs_decade["efficacy_tier"] = ms_dmt_procs_decade["drug_generic_name"].map(tier_map)
ms_dmt_procs_decade["source_table"] = "Procedures"

ms_dmt_decade = pd.concat([ms_dmt_meds_decade, ms_dmt_procs_decade], ignore_index=True, sort=False)
log(f"ms_dmt_decade: {len(ms_dmt_decade):,} rows, "
    f"{ms_dmt_decade['ndid'].nunique():,} treated patients (of {total_cohort_n:,})")

# ===== cell 7 =====
# ── Shared helper: build the full battery of analyses for a given ms_dmt table ─────
def build_analysis_bundle(ms_dmt, label):
    log(f"building analysis bundle: {label}")
    n_dmt_treated = ms_dmt["ndid"].nunique()

    dmt_overall_utilization = pd.DataFrame([
        ("Total MS patients", total_cohort_n, None),
        ("Patients receiving >=1 DMT", n_dmt_treated, round(100 * n_dmt_treated / total_cohort_n, 1)),
    ], columns=["Characteristic", "N", "%"])

    class_patients = ms_dmt.groupby("drug_class")["ndid"].nunique()
    dmt_utilization_by_class = pd.DataFrame({
        "N_patients": class_patients,
        "% of total cohort": (100 * class_patients / total_cohort_n).round(1),
    }).sort_values("N_patients", ascending=False)

    drug_patients = ms_dmt.groupby("drug_generic_name")["ndid"].nunique()
    dmt_utilization_by_drug = pd.DataFrame({
        "N_patients": drug_patients,
        "% of DMT-treated patients": (100 * drug_patients / n_dmt_treated).round(1),
    }).sort_values("N_patients", ascending=False)
    dmt_utilization_by_drug.insert(0, "Rank", range(1, len(dmt_utilization_by_drug) + 1))

    # first-DMT efficacy tier + class annotation
    first_dmt_record = ms_dmt.sort_values("med_start_date").groupby("ndid").first()
    n_treated_patients = len(first_dmt_record)
    tier_counts = first_dmt_record["efficacy_tier"].value_counts()
    first_dmt_tier_distribution = pd.DataFrame([
        ("N (DMT-treated patients)", n_treated_patients, None),
        ("Standard efficacy", int(tier_counts.get("Standard efficacy", 0)),
         round(100 * tier_counts.get("Standard efficacy", 0) / n_treated_patients, 1)),
        ("High efficacy", int(tier_counts.get("High efficacy", 0)),
         round(100 * tier_counts.get("High efficacy", 0) / n_treated_patients, 1)),
    ], columns=["Metric", "N", "%"])

    _class_counts = first_dmt_record.groupby(["efficacy_tier", "drug_class"]).size().rename("N").reset_index()
    first_dmt_tier_by_class = _class_counts.copy()
    first_dmt_tier_by_class["%_of_tier"] = first_dmt_tier_by_class.apply(
        lambda r: round(100 * r["N"] / tier_counts.get(r["efficacy_tier"], 1), 1), axis=1)
    first_dmt_tier_by_class = first_dmt_tier_by_class.sort_values(
        ["efficacy_tier", "N"], ascending=[True, False]).reset_index(drop=True)

    # time-to-first-DMT, true first-ever diagnosis date -> first DMT within this table
    first_dmt_per_patient = ms_dmt.groupby("ndid")["med_start_date"].min().rename("first_dmt_date")
    ttd = ms_cohort_demo[["ndid", "first_diagnosis_date", "rrms_flag", "ppms_flag", "spms_flag",
                           "ms_unspecified_flag"]].merge(first_dmt_per_patient, on="ndid", how="inner")
    ttd["first_diagnosis_date"] = pd.to_datetime(ttd["first_diagnosis_date"], errors="coerce")
    ttd["days_to_first_dmt"] = (ttd["first_dmt_date"] - ttd["first_diagnosis_date"]).dt.days

    def _window_v2(d):
        if pd.isna(d): return "Unknown"
        if d <= 30: return "<=30 days"
        if d <= 90: return "31-90 days"
        if d <= 365: return "91-365 days"
        return ">365 days"

    ttd["initiation_window"] = ttd["days_to_first_dmt"].apply(_window_v2)
    n_treated_ttd = len(ttd)
    window_order = ["<=30 days", "31-90 days", "91-365 days", ">365 days", "Unknown"]
    win_counts = ttd["initiation_window"].value_counts()
    time_to_first_dmt_windows = pd.DataFrame(
        [("N (DMT-treated patients, resolvable)", n_treated_ttd, None, None),
         ("Median days to first DMT", None, None, round(ttd["days_to_first_dmt"].median(), 1)),
         ("Mean days to first DMT", None, None, round(ttd["days_to_first_dmt"].mean(), 1))] +
        [(w, int(win_counts.get(w, 0)), round(100 * win_counts.get(w, 0) / n_treated_ttd, 1), None)
         for w in window_order if win_counts.get(w, 0) > 0],
        columns=["Metric", "N", "%", "Value"])

    # tier switching matrix
    patient_drug_first_start = (
        ms_dmt.groupby(["ndid", "drug_generic_name"])["med_start_date"].min()
        .reset_index().sort_values(["ndid", "med_start_date"]))
    patient_drug_first_start["efficacy_tier"] = patient_drug_first_start["drug_generic_name"].map(tier_map)
    _seq = patient_drug_first_start.groupby("ndid").agg(
        drug_list=("drug_generic_name", list), tier_list=("efficacy_tier", list),
        date_list=("med_start_date", list))
    _seq["n_drugs"] = _seq["drug_list"].str.len()
    multi_dmt_seq = _seq[_seq["n_drugs"] > 1].copy()
    n_multi_dmt_patients = len(multi_dmt_seq)

    tier_transition_rows = []
    for ndid, row in multi_dmt_seq.iterrows():
        tiers, dates = row["tier_list"], row["date_list"]
        for step, (from_tier, to_tier, from_date, to_date) in enumerate(
            zip(tiers[:-1], tiers[1:], dates[:-1], dates[1:]), start=1):
            tier_transition_rows.append((ndid, step, from_tier, to_tier, from_date, to_date))
    tier_transitions_long = pd.DataFrame(
        tier_transition_rows, columns=["ndid", "step", "from_tier", "to_tier", "from_date", "to_date"])
    tier_transitions_long["gap_days"] = (tier_transitions_long["to_date"] - tier_transitions_long["from_date"]).dt.days
    tier_transitions_long["tier_transition"] = tier_transitions_long["from_tier"] + " -> " + tier_transitions_long["to_tier"]
    tier_transitions_long["switch_year"] = tier_transitions_long["to_date"].dt.year

    total_transitions = len(tier_transitions_long)
    tt_counts = tier_transitions_long["tier_transition"].value_counts()
    _tier_labels = [
        ("Standard efficacy -> Standard efficacy", "Within-tier (Standard -> Standard)"),
        ("High efficacy -> High efficacy",         "Within-tier (High -> High)"),
        ("Standard efficacy -> High efficacy",     "Cross-tier escalation (Standard -> High)"),
        ("High efficacy -> Standard efficacy",     "Cross-tier de-escalation (High -> Standard)"),
    ]
    dmt_tier_switch_matrix = pd.DataFrame(
        [(label, int(tt_counts.get(key, 0)),
          round(100 * tt_counts.get(key, 0) / total_transitions, 1) if total_transitions else None,
          round(100 * tt_counts.get(key, 0) / n_multi_dmt_patients, 1) if n_multi_dmt_patients else None)
         for key, label in _tier_labels],
        columns=["Transition type", "N_transitions", "% of total transitions", "% of multi-DMT patients"])

    n_switches_per_patient = (multi_dmt_seq["n_drugs"] - 1).reindex(ms_cohort_demo["ndid"], fill_value=0)

    def _switch_bucket(v):
        if v == 0: return "0 switches"
        if v == 1: return "1 switch"
        if v == 2: return "2 switches"
        if v == 3: return "3 switches"
        return "4+ switches"

    switch_count_dist = n_switches_per_patient.apply(_switch_bucket).value_counts()
    switch_bucket_order = ["0 switches", "1 switch", "2 switches", "3 switches", "4+ switches"]
    dmt_switch_frequency = pd.DataFrame(
        [(b, int(switch_count_dist.get(b, 0)), round(100 * switch_count_dist.get(b, 0) / total_cohort_n, 1))
         for b in switch_bucket_order], columns=["N_switches", "N_patients", "% of full cohort"])

    _gap = tier_transitions_long["gap_days"].dropna()
    n_escalation = int(tt_counts.get("Standard efficacy -> High efficacy", 0))
    n_deescalation = int(tt_counts.get("High efficacy -> Standard efficacy", 0))
    n_cross_tier = n_escalation + n_deescalation
    switch_velocity_summary = pd.DataFrame([
        ("N individual switches", total_transitions, None),
        ("N multi-DMT patients (>=1 switch)", n_multi_dmt_patients,
         round(100 * n_multi_dmt_patients / n_dmt_treated, 1)),
        ("Median days between consecutive switches", None, round(_gap.median(), 1)),
        ("Mean days between consecutive switches", None, round(_gap.mean(), 1)),
        ("Cross-tier escalation (Standard->High)", n_escalation,
         round(100 * n_escalation / total_transitions, 1) if total_transitions else None),
        ("Cross-tier de-escalation (High->Standard)", n_deescalation,
         round(100 * n_deescalation / total_transitions, 1) if total_transitions else None),
        ("Escalation as % of all cross-tier transitions", None,
         round(100 * n_escalation / n_cross_tier, 1) if n_cross_tier else None),
    ], columns=["Metric", "N", "%/Value"])

    dmt_courses = ms_dmt.groupby(["ndid", "drug_generic_name", "efficacy_tier"]).agg(
        discont_date=("discont_date", "max"),
        discont_reason=("discont_reason", lambda s: s.dropna().iloc[0] if s.notna().any() else None),
    ).reset_index()
    dmt_courses["discontinued"] = dmt_courses["discont_date"].notna()
    _discont = dmt_courses[dmt_courses["discontinued"]]
    _reason_counts = _discont["discont_reason"].fillna("Unknown / not recorded").value_counts()
    dmt_switch_reasons = pd.DataFrame({
        "N": _reason_counts,
        "% of discontinued courses": (100 * _reason_counts / len(_discont)).round(1) if len(_discont) else 0,
    }).reset_index(names="reason")
    dmt_switch_reasons.insert(0, "Rank", range(1, len(dmt_switch_reasons) + 1))

    return dict(
        n_dmt_treated=n_dmt_treated,
        dmt_overall_utilization=dmt_overall_utilization,
        dmt_utilization_by_class=dmt_utilization_by_class,
        dmt_utilization_by_drug=dmt_utilization_by_drug,
        first_dmt_tier_distribution=first_dmt_tier_distribution,
        first_dmt_tier_by_class=first_dmt_tier_by_class,
        time_to_first_dmt_windows=time_to_first_dmt_windows,
        dmt_tier_switch_matrix=dmt_tier_switch_matrix,
        dmt_switch_frequency=dmt_switch_frequency,
        switch_velocity_summary=switch_velocity_summary,
        dmt_switch_reasons=dmt_switch_reasons,
        tier_transitions_long=tier_transitions_long,
    )


uncapped = build_analysis_bundle(ms_dmt_uncapped, "Primary (Uncapped)")
decade = build_analysis_bundle(ms_dmt_decade, "Secondary (Decade 2015-2026-06)")

# ===== cell 8 =====
# ── NEW: decade medication transition TRENDS -- year-by-year tier-switch counts ────
_tl = decade["tier_transitions_long"]
_tl_valid = _tl[_tl["switch_year"].between(2015, 2026)]
decade_switch_trend_by_year = (
    _tl_valid.groupby(["switch_year", "tier_transition"]).size().unstack(fill_value=0)
)
_rename = {
    "Standard efficacy -> Standard efficacy": "Standard -> Standard",
    "High efficacy -> High efficacy": "High -> High",
    "Standard efficacy -> High efficacy": "Standard -> High (escalation)",
    "High efficacy -> Standard efficacy": "High -> Standard (de-escalation)",
}
decade_switch_trend_by_year = decade_switch_trend_by_year.rename(columns=_rename)
for col in _rename.values():
    if col not in decade_switch_trend_by_year.columns:
        decade_switch_trend_by_year[col] = 0
decade_switch_trend_by_year = decade_switch_trend_by_year[list(_rename.values())]
decade_switch_trend_by_year["Total switches"] = decade_switch_trend_by_year.sum(axis=1)
decade_switch_trend_by_year["% escalation of cross-tier"] = (
    100 * decade_switch_trend_by_year["Standard -> High (escalation)"] /
    (decade_switch_trend_by_year["Standard -> High (escalation)"] + decade_switch_trend_by_year["High -> Standard (de-escalation)"]).replace(0, np.nan)
).round(1)
log("built decade_switch_trend_by_year")
decade_switch_trend_by_year

# ===== cell 9 =====
# ── Cohort characteristics (age-cleaned) ─────────────────────────────────────────────
n = total_cohort_n
age = ms_cohort_demo["age_at_diagnosis"]
onset_counts = ms_cohort_demo["onset_tier_v2"].value_counts()
onset_order = ["Pediatric-Onset (<18)", "Unclassified (18)", "Adult-Onset (19-49)", "Late-Onset (50+)", "Unknown"]
sex_counts = ms_cohort_demo["gender_bucket"].value_counts()

cohort_characteristics_clean = pd.DataFrame(
    [
        ("Total MS patients (age >=100 excluded)", n, None, None),
        ("Patients excluded for age >=100 at diagnosis", len(excluded_age100), None, None),
        ("Mean age at diagnosis (true, full history)", None, None, round(age.mean(), 1)),
        ("SD", None, None, round(age.std(), 1)),
        ("Median", None, None, round(age.median(), 0)),
        ("Q1 (25th pct)", None, None, round(age.quantile(0.25), 0)),
        ("Q3 (75th pct)", None, None, round(age.quantile(0.75), 0)),
        ("Max age at diagnosis (post-cleanup)", None, None, round(age.max(), 0)),
    ] + [
        (label, int(sex_counts.get(label, 0)), round(100 * sex_counts.get(label, 0) / n, 1), None)
        for label in ["Female", "Male", "Unknown"]
    ] + [
        (label, int(onset_counts.get(label, 0)), round(100 * onset_counts.get(label, 0) / n, 1), None)
        for label in onset_order if onset_counts.get(label, 0) > 0
    ] + [
        (label, int(ms_cohort_demo[flag].sum()), round(100 * ms_cohort_demo[flag].sum() / n, 1), None)
        for label, flag in [("RRMS", "rrms_flag"), ("PPMS", "ppms_flag"), ("SPMS", "spms_flag"),
                             ("Unspecified", "ms_unspecified_flag")]
    ],
    columns=["Characteristic", "N", "%", "Value"],
)
log("built cohort_characteristics_clean")
cohort_characteristics_clean

# ===== cell 10 =====
# ── Write consolidated workbook ──────────────────────────────────────────────────────
OUT_PATH = "/Users/anjali28/Desktop/MS demo paper/ms_summary_tables5.xlsx"

SHEET_MARK = "— sheet overview —"
readme_rows = [
    ("cohort_characteristics_clean", SHEET_MARK, "Age-cleaned (age>=100 excluded) full-cohort demographics: age, sex, onset tier, subtype."),
    ("uncapped_dmt_overall_util", SHEET_MARK, "Primary Analysis (no date restriction): % of cohort ever receiving >=1 DMT."),
    ("uncapped_dmt_util_by_class", SHEET_MARK, "Primary Analysis: unique patients per drug class / full cohort."),
    ("uncapped_dmt_util_by_drug", SHEET_MARK, "Primary Analysis: unique patients per individual DMT / DMT-treated patients."),
    ("uncapped_first_dmt_tier_dist", SHEET_MARK, "Primary Analysis: Standard vs. High efficacy tier of each patient's first-ever DMT record."),
    ("uncapped_first_dmt_tier_cls", SHEET_MARK, "Primary Analysis: drug class of first DMT, within each efficacy tier."),
    ("uncapped_ttd_windows", SHEET_MARK, "Primary Analysis: days from true first diagnosis to first-ever DMT record, no date cap, 4 mutually exclusive windows."),
    ("uncapped_tier_switch_matrix", SHEET_MARK, "Primary Analysis: all drug-switch events, full history, classified by efficacy-tier transition type."),
    ("uncapped_switch_frequency", SHEET_MARK, "Primary Analysis: number of distinct-drug switches per patient, full history."),
    ("uncapped_switch_velocity", SHEET_MARK, "Primary Analysis: switch velocity and escalation/de-escalation counts, full history."),
    ("uncapped_switch_reasons", SHEET_MARK, "Primary Analysis: discontinuation reasons ranked, full history."),
    ("decade_dmt_overall_util", SHEET_MARK, "Secondary Analysis (2015-2026-06 only): % of cohort ever receiving >=1 DMT within the decade window."),
    ("decade_dmt_util_by_class", SHEET_MARK, "Secondary Analysis: unique patients per drug class / full cohort, decade window."),
    ("decade_dmt_util_by_drug", SHEET_MARK, "Secondary Analysis: unique patients per individual DMT / DMT-treated patients, decade window."),
    ("decade_first_dmt_tier_dist", SHEET_MARK, "Secondary Analysis: Standard vs. High efficacy tier of first DMT record within the decade window."),
    ("decade_first_dmt_tier_cls", SHEET_MARK, "Secondary Analysis: drug class of first DMT within the decade window, by tier."),
    ("decade_ttd_windows", SHEET_MARK, "Secondary Analysis: days from true first diagnosis to first DMT record within the decade window, 4 mutually exclusive windows."),
    ("decade_tier_switch_matrix", SHEET_MARK, "Secondary Analysis: drug-switch events within the decade window, classified by efficacy-tier transition type."),
    ("decade_switch_frequency", SHEET_MARK, "Secondary Analysis: number of distinct-drug switches per patient, decade window."),
    ("decade_switch_velocity", SHEET_MARK, "Secondary Analysis: switch velocity and escalation/de-escalation counts, decade window."),
    ("decade_switch_reasons", SHEET_MARK, "Secondary Analysis: discontinuation reasons ranked, decade window."),
    ("decade_switch_trend_by_year", SHEET_MARK, "NEW: year-by-year count of each efficacy-tier switch-transition type within 2015-2026-06, showing whether escalation vs. de-escalation prevalence has shifted over the decade."),
    ("Methodology", "Age cleanup", f"Patients with age at diagnosis >= 100 years excluded from the cohort entirely ({len(excluded_age100)} patients removed, {n_before:,} -> {n_after:,})."),
    ("Methodology", "Primary Analysis (Uncapped)", "No date restriction on DMT records (Medications + Procedures, full history)."),
    ("Methodology", "Secondary Analysis (Decade)", f"DMT records (Medications + Procedures) restricted to {DECADE_START}..{DECADE_END}."),
]
readme = pd.DataFrame(readme_rows, columns=["Sheet", "Column", "Description"])

with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
    readme.to_excel(writer, sheet_name="README", index=False)
    cohort_characteristics_clean.to_excel(writer, sheet_name="cohort_characteristics_clean", index=False)
    for prefix, bundle in [("uncapped", uncapped), ("decade", decade)]:
        bundle["dmt_overall_utilization"].to_excel(writer, sheet_name=f"{prefix}_dmt_overall_util", index=False)
        bundle["dmt_utilization_by_class"].to_excel(writer, sheet_name=f"{prefix}_dmt_util_by_class")
        bundle["dmt_utilization_by_drug"].to_excel(writer, sheet_name=f"{prefix}_dmt_util_by_drug")
        bundle["first_dmt_tier_distribution"].to_excel(writer, sheet_name=f"{prefix}_first_dmt_tier_dist", index=False)
        bundle["first_dmt_tier_by_class"].to_excel(writer, sheet_name=f"{prefix}_first_dmt_tier_cls", index=False)
        bundle["time_to_first_dmt_windows"].to_excel(writer, sheet_name=f"{prefix}_ttd_windows", index=False)
        bundle["dmt_tier_switch_matrix"].to_excel(writer, sheet_name=f"{prefix}_tier_switch_matrix", index=False)
        bundle["dmt_switch_frequency"].to_excel(writer, sheet_name=f"{prefix}_switch_frequency", index=False)
        bundle["switch_velocity_summary"].to_excel(writer, sheet_name=f"{prefix}_switch_velocity", index=False)
        bundle["dmt_switch_reasons"].to_excel(writer, sheet_name=f"{prefix}_switch_reasons", index=False)
    decade_switch_trend_by_year.to_excel(writer, sheet_name="decade_switch_trend_by_year")

print(f"Wrote {OUT_PATH}")

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FILL = PatternFill(start_color="FF305496", end_color="FF305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
sheet_descriptions = {s: d for s, c, d in readme_rows if c == SHEET_MARK}

wb = load_workbook(OUT_PATH)
for name in wb.sheetnames:
    if name in ("README", "Methodology"):
        continue
    ws = wb[name]
    desc = sheet_descriptions.get(name, "")
    max_col = max(ws.max_column, 1)
    ws.insert_rows(1)
    ws.cell(row=1, column=1, value=f"{name} -- {desc}" if desc else name)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws["A1"].fill = HEADER_FILL
    ws["A1"].font = HEADER_FONT
    ws["A1"].alignment = WRAP_TOP
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = ws.cell(row=3, column=1)
wb.save(OUT_PATH)
print(f"Added heading banners; re-saved {OUT_PATH}")
